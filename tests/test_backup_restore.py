"""Tümünü indir (zip) → geri yükle; şifreli yedek; Excel; zamanlanmış işler her kullanıcı alanında."""
import io
import os
import zipfile

from openpyxl import load_workbook
from tests.conftest import fixture_bytes


def _entry(client, desc, amount=100):
    return client.post("/api/entries", json={"entry_date": "2026-10-01", "description": desc, "amount": amount,
                                             "kind": "gider", "category": "Market & Gıda", "payment": "nakit"}).json()["id"]


def test_export_zip_and_restore(client, tmpdb):
    client.post("/api/statements/upload", files=[("files", ("a.pdf", fixture_bytes("vakifbank_a.txt"), "application/pdf"))])
    _entry(client, "Ekmek")
    z = client.get("/api/settings/export.zip")
    assert z.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(z.content)).namelist()
    assert "hesapkitap.db" in names and "OKUBENI.txt" in names and any(n.startswith("uploads/") for n in names)

    # sonra veri bozuldu: kayıt silindi, ekstre PDF'i kayboldu
    with tmpdb.connect() as c:
        c.execute("DELETE FROM entries")
    for f in os.listdir(tmpdb.upload_dir()):
        os.remove(os.path.join(tmpdb.upload_dir(), f))
    r = client.post("/api/settings/backup/restore", files={"file": ("yedek.zip", z.content, "application/zip")}).json()
    assert r["entries"] == 1 and r["statements"] == 1 and r["files"] >= 1 and "geri-yukleme-oncesi" in r["before"]
    assert [e["description"] for e in tmpdb.list_entries()] == ["Ekmek"]
    assert os.listdir(tmpdb.upload_dir())
    # geri yükleme öncesi hal listede: oraya da dönülebilir
    assert client.post(f"/api/settings/backup/{r['before']}/restore").json()["entries"] == 0
    assert tmpdb.list_entries() == []


def test_restore_rejects_garbage_and_needs_passphrase(client, tmpdb):
    from app import backup

    bad = client.post("/api/settings/backup/restore", files={"file": ("x.db", b"merhaba", "application/octet-stream")})
    assert bad.status_code == 422
    _entry(client, "Şifreli")
    path = backup.make_backup()
    blob = backup.encrypt(open(path, "rb").read(), "gizli-söz")
    with tmpdb.connect() as c:
        c.execute("DELETE FROM entries")
    f = {"file": ("ekstre.db.enc", blob, "application/octet-stream")}
    assert "şifre" in client.post("/api/settings/backup/restore", files=f).json()["detail"]
    assert client.post("/api/settings/backup/restore", files=f, data={"passphrase": "yanlış"}).status_code == 422
    assert client.post("/api/settings/backup/restore", files=f, data={"passphrase": "gizli-söz"}).json()["entries"] == 1


def test_export_xlsx(client, tmpdb):
    _entry(client, "Peynir", 250)
    wb = load_workbook(io.BytesIO(client.get("/api/settings/export.xlsx").content))
    assert wb.sheetnames == ["İşlemler", "Sabit kalemler", "Hesaplar", "Borç-alacak"]
    rows = list(wb["İşlemler"].values)
    assert rows[1][2] == "Peynir" and rows[1][5] == -250


def test_jobs_run_in_each_user_space(tmpdb, tmp_path, monkeypatch):
    from datetime import datetime

    from app import jobs, system

    u = system.create_user("esim", "x" * 60, role="user")
    spaces = jobs.spaces()
    assert spaces[0] == (None, None) and any(s[1] and s[1]["id"] == u["id"] for s in spaces)
    with tmpdb.use_space(system.user_db_path(u), u):
        names = [n for n, _ in jobs.due_jobs(datetime(2026, 10, 26, 23, 0))]
        assert "backup" in names and "remind" not in names  # Telegram'a yazan iş başkasının alanında çalışmaz
    names = [n for n, _ in jobs.due_jobs(datetime(2026, 10, 26, 23, 0))]
    assert "remind" in names
