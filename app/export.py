"""Excel'e aktarma: işlemler, sabit kalemler, hesaplar, borç/alacak. Başka bir programa taşımak ya da kendi tablonu
yapmak için; geri yükleme için değil (onun için zip yedek).
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app import db, finance

KIND = {"harcama": "Kart harcaması", "iade": "İade", "faiz_ucret": "Faiz/ücret", "odeme": "Kart ödemesi",
        "gelir": "Gelir", "gider": "Gider", "borc": "Borç hareketi"}


def _sheet(wb: Workbook, title: str, head: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, h in enumerate(head, 1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:500]])
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
    for i, h in enumerate(head, 1):
        if h.startswith("Tutar") or h in ("Bakiye", "Kalan"):
            for c in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for cell in c:
                    cell.number_format = "#,##0.00"


def workbook() -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    rows = [r for r in finance.transactions(None)
            if not r.get("pending") and not r.get("deleted") and not r.get("linked") and r.get("source") != "estimate"]
    rows.sort(key=lambda r: r["date"])
    out_kinds = {"harcama", "faiz_ucret", "gider"}
    _sheet(wb, "İşlemler", ["Tarih", "Dönem", "Açıklama", "Tür", "Kategori", "Tutar (₺)", "Kart", "Hesap", "Kişi", "Taksit", "Not"],
           [[r["date"][:10], r["month"], r["description"], KIND.get(r["kind"], r["kind"]), r["category"],
             round(-abs(r["amount"]) if r["kind"] in out_kinds else r["amount"], 2), r.get("card"), r.get("account") or r.get("bank"),
             r.get("counterparty"), r.get("installment"), r.get("note")] for r in rows])
    _sheet(wb, "Sabit kalemler", ["Açıklama", "Tür", "Kategori", "Tutar (₺)", "Gün", "Başlangıç", "Bitiş"],
           [[r["description"], r["kind"], r["category"], r["amount"], r["day_of_month"], r["start_month"], r["end_month"]]
            for r in db.list_recurring()])
    _sheet(wb, "Hesaplar", ["Hesap", "Tür", "Bakiye", "Tarih"],
           [[a["name"], a["kind"], a["current"] if a["current"] is not None else a["balance"], a["as_of"]] for a in db.list_accounts()])
    _sheet(wb, "Borç-alacak", ["Kişi", "Yön", "Birim", "Tutar", "Kalan", "Vade", "Not"],
           [[d["person"], "Alacak" if d["direction"] == "alacak" else "Borç", d["asset"], d["amount"], d.get("remaining"),
             d.get("due_date"), d.get("note")] for d in db.list_debts()])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
