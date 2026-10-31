"""Banka hesap hareketleri (vadesiz hesap) içe aktarma: Excel (.xlsx/.xls) ya da CSV.

Bankalar farklı biçim kullanır; sütunlar başlıklardan tahmin edilir, kullanıcı önizlemede düzeltebilir.
- Gelen para → gelir, giden para → gider (banka kartı / havale)
- Kredi kartı borç ödemesi → gider sayılmaz (harcamalar ekstreden gelir), ilgili ekstre "ödendi" yapılır
- Kendi hesapları arası virman → atlanır
- Maaş → maaş sabit kalemine bağlanır (onay bekleyenden düşer)
- Aynı hareket ikinci kez eklenmez
"""
import csv
import hashlib
import io
import re
from datetime import date, datetime

from app.textnorm import normalize

HEADER_HINTS = {
    "date": ("TARIH", "ISLEM TARIHI", "VALOR", "DATE"),
    "description": ("ACIKLAMA", "ISLEM ACIKLAMASI", "ISLEM", "AYRINTI", "DETAY", "DESCRIPTION"),
    "amount": ("TUTAR", "ISLEM TUTARI", "AMOUNT"),
    "debit": ("BORC", "CIKAN", "GIDEN", "HARCAMA"),
    "credit": ("ALACAK", "GIREN", "GELEN", "YATAN"),
    "balance": ("BAKIYE", "BALANCE"),
}
CARD_PAYMENT = re.compile(r"KREDI KARTI|KART BORC|KK BORC|KK ODEME|KARTI ODEME|KART ODEME|CREDIT CARD")
TRANSFER = re.compile(r"HESAPLAR ARASI|VIRMAN|KENDI HESAB|HESABIMA|ARA HESAP|DOVIZ ALIS|DOVIZ SATIS|VADELI|KATILMA HESAB")
SALARY = re.compile(r"MAAS|UCRET ODEME|PERSONEL ODEME|BORDRO")


class ImportError_(ValueError):
    pass


# ---------------------------------------------------------------- dosya → satırlar

def read_rows(data: bytes, filename: str) -> list[list[str]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    if name.endswith(".xls"):
        import xlrd

        book = xlrd.open_workbook(file_contents=data)
        sh = book.sheet_by_index(0)
        out = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode).date().isoformat())
                else:
                    row.append(_cell(cell.value))
            out.append(row)
        return out
    # CSV: kodlama ve ayraç tahmini
    for enc in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportError_("Dosya okunamadı")
    try:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=";,\t|")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if text.count(";") > text.count(",") else ","
    return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (datetime, date)):
        return (v.date() if isinstance(v, datetime) else v).isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------- sütun tahmini

def detect(rows: list[list[str]]) -> dict:
    """Başlık satırını ve sütun eşlemesini tahmin et."""
    best, best_score = 0, -1
    for i, row in enumerate(rows[:40]):
        norm = [normalize(c) for c in row]
        score = sum(any(any(h == n or n.startswith(h) for h in hints) for n in norm) for hints in HEADER_HINTS.values())
        if score > best_score:
            best, best_score = i, score
    header = [normalize(c) for c in rows[best]] if rows else []
    mapping: dict[str, int | None] = {}
    for key, hints in HEADER_HINTS.items():
        idx = next((j for j, n in enumerate(header) if any(n == h or n.startswith(h) for h in hints)), None)
        # "ISLEM TARIHI" hem tarih hem açıklama ipucuna uymasın: tarih sütunu açıklama olamaz
        if key == "description" and idx is not None and idx == mapping.get("date"):
            idx = next((j for j, n in enumerate(header) if j != idx and any(n.startswith(h) for h in hints)), None)
        mapping[key] = idx
    return {"header_row": best, "columns": rows[best] if rows else [], "mapping": mapping}


# ---------------------------------------------------------------- değer ayrıştırma

def parse_amount(v: str) -> float | None:
    t = re.sub(r"[^\d,.\-+]", "", v or "")
    if not t or t in "-+":
        return None
    neg = t.startswith("-") or (v or "").strip().endswith("-")
    t = t.lstrip("+-").rstrip("-")
    if "," in t and "." in t:
        t = t.replace(".", "").replace(",", ".") if t.rfind(",") > t.rfind(".") else t.replace(",", "")
    elif "," in t:
        t = t.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", t):
        t = t.replace(".", "")
    try:
        x = float(t)
    except ValueError:
        return None
    return -x if neg else x


def parse_date(v: str) -> date | None:
    v = (v or "").strip()[:19]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y %H:%M", "%d.%m.%Y %H:%M:%S",
                "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S", "%d.%m.%y"):
        try:
            return datetime.strptime(v, fmt).date()
        except ValueError:
            continue
    return None


def _squash(s: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", normalize(s))


def _parties(desc: str) -> tuple[str, str] | None:
    """Açıklamadaki (gönderen, alıcı) ham metni (normalize edilmiş), kalıp tanınmazsa None."""
    n = normalize(desc)
    pats = (
        # VakıfBank: "… sorgu no'lu [Banka] X hesabından [Banka] Y hesabına giden/gelen"
        r"NO'?\s?LU\s+(.+?)\s+HESABINDAN\s+(.+?)\s+HESABINA",
        # Halkbank giden: "GIDEN FAST - X TARAFINDAN GONDERILEN,Y,Banka…" / HAVALE: "…,X,TR..,Y,TR.."
        r"-\s*([A-Z ]{5,}?)\s+TARAFINDAN GONDERILEN\s*,\s*(?:[^,]+?,\s*TR[\d* ]+,\s*)?([^,]+?)\s*,",
        # VakıfBank gelen EFT: "… sorgu numaralı ORNEK SIRKET A.S. tarafından AYSE YILMAZ tarafına gelen EFT"
        r"SORGU NUMARALI\s+(.+?)\s+TARAFINDAN\s+(.+?)\s+TARAFINA",
        # Halkbank havale (açıklamalı): "not,X,TRxx…,Y,TRxx…"
        r"(?:^|,)\s*([^,]+?)\s*,\s*TR\d[\d ]{10,}\s*,\s*([^,]+?)\s*,\s*TR\d",
    )
    for pat in pats:
        m = re.search(pat, n)
        if m and _clean_name(m.group(1)) and _clean_name(m.group(2)):
            return m.group(1), m.group(2)
    # Halkbank gelen: "not,X 'DAN GELEN FAST ODEMESI" — alıcı yazılmaz; kendi hesabından gelende not alanı gönderenin
    # adıdır (otomatik): (gönderen, not) döner
    m = re.search(r"^\s*([^,]{3,}?)\s*,\s*([A-Z ]{5,}?)\s*'\s*D\s*[AE]\s*N\s+GELEN", n)
    if m and _clean_name(m.group(2)):
        return m.group(2), m.group(1)
    return None


def _sender_receiver(desc: str) -> tuple[str, str] | None:
    """Gönderen ve alıcı adı (banka adları atılmış, boşluksuz), kalıp tanınmazsa None. Hesap sahibinin adı bilinmese de
    (kural tabanlı okuyucular adı vermez, kart ekstresi henüz yok) kendi hesapları arası aktarımı tanır."""
    p = _parties(desc)
    if p:
        return _squash(_clean_name(p[0])), _squash(_clean_name(p[1]))
    sq = _squash(desc)
    m = re.search(r"GONDEREN([A-Z]+?)ALICI([A-Z]+?)(FAST|EFT|HAVALE|PARA|$)", sq)
    return (m.group(1), m.group(2)) if m else None


def is_self_transfer(desc: str, owners: list[str]) -> bool:
    """Kendi hesapları arası: GÖNDEREN ve ALICI aynı kişi (boşluksuz yazımlar dahil, ör. Kuveyt Türk). Kullanıcının adı
    yalnız serbest açıklama metninde geçiyorsa ('Ayşe Yılmaz kurs ücreti … X hesabından Can Ö. hesabına')
    virman DEĞİLDİR: kalıp tanınıyorsa karar gönderen/alıcıya göre verilir."""
    return classify_reason(desc, -1, owners)[0] == "virman"


# Kurum/şirket (emeklilik, sigorta, ödeme kuruluşu): kendi hesapları arası sayılmaz (kullanıcının takip ettiği hesap değil)
_INSTITUTION = re.compile(r"\b(EMEKLILIK|SIGORTA|HAYAT|PAYMENT|PAYMENTS|CORP|INC|LLC|LTD|LIMITED|SIRKETI|A\.?S\.? TARAFINDAN|"
                          r"FON|PORTFOY|YATIRIM|MENKUL)\b")


def classify_reason(desc: str, amount: float, owners: list[str] | None = None) -> tuple[str, str]:
    """(tür, neden) — neden onarım raporunda gösterilir."""
    n = normalize(desc)
    sq = _squash(desc)
    if amount < 0 and (CARD_PAYMENT.search(n) or re.search(r"KREDIKARTI(BORC|ODEME)|KARTBORC|KKBORC", sq)):
        return "kart_odemesi", "kart borcu ödemesi"
    plain = "gelir" if amount > 0 else "gider"
    sr = _sender_receiver(desc)
    if sr and len(sr[0]) >= 6:
        if sr[0] == sr[1]:
            return "virman", "gönderen = alıcı"
        if not re.search(r"HESAPLAR ARASI|VIRMAN|KENDI HESAB", n):
            return plain, "gönderen ≠ alıcı"  # açıklamada geçen ad ya da 'VADELI' gibi kelime karar vermez
    if _INSTITUTION.search(n) and not re.search(r"HESAPLAR ARASI|VIRMAN|KENDI HESAB", n):
        return plain, "kurum/şirket"
    if re.search(r"PARA CEK|CEKILEN\b|\bPR CKM\b", n) and not re.search(r"HESAPLAR ARASI|VIRMAN", n):
        return plain, "para çekme"  # 'Mevduat Para Çek … AYŞE YILMAZ ne ait … hesabından çekilen': nakit, virman değil
    if m := TRANSFER.search(n):
        return "virman", f"açıklamada '{m.group(0)}'"
    names = {_squash(o) for o in owners or [] if o and len(_squash(o)) >= 6}
    for nm in names:  # kalıp tanınmadı: "X hesabından … X hesabına" (adı iki kez geçiyor)
        if sq.count(nm) >= 2 and re.search(r"HESABINDAN|TARAFINDANGONDERILEN|GONDEREN", sq):
            return "virman", "kendi adı iki kez (gönderen ve alıcı)"
    return plain, ""


def classify(desc: str, amount: float, owners: list[str] | None = None) -> str:
    return classify_reason(desc, amount, owners)[0]


def finalize(rows: list[dict], owners: list[str]) -> list[dict]:
    """PDF'ten gelen satırlara tür ve parmak izi ekle."""
    out = []
    for r in rows:
        d = parse_date(r["date"])
        if not d or not r.get("amount"):
            continue
        # Parmak izi açıklamaya bakmaz: model aynı satırı her okumada biraz farklı yazabilir; tarih + tutar +
        # bakiye (zincirle doğrulanmış) hareketi tek başına tanımlar, aynı PDF tekrar yüklenince çift kayıt olmaz
        bal = r.get("balance")
        out.append({"date": d.isoformat(), "description": r["description"][:500], "amount": round(r["amount"], 2),
                    "type": classify(r["description"], r["amount"], owners), "balance": bal,
                    "hash": row_hash(d, "" if bal is not None else r["description"], r["amount"], f"{bal:.2f}" if bal is not None else "")})
    _unique_hashes(out)
    pair_cancellations(out)
    return out


def _unique_hashes(rows: list[dict]) -> None:
    """Aynı gün, aynı tutar, aynı bakiye iki ayrı hareket olabilir (çekim → iptal → yine aynı çekim): parmak izi aynı
    çıkar ve ikincisi 'zaten var' sanılıp kaybolurdu. Belgedeki 2., 3. tekrar sıra numarasıyla ayrışır; ilki eskisiyle
    aynı kalır (önceden yüklenmiş kayıtlar yeniden yüklemede kopyalanmaz)."""
    seen: dict[str, int] = {}
    for r in rows:
        n = seen.get(r["hash"], 0)
        seen[r["hash"]] = n + 1
        if n:
            r["hash"] = hashlib.sha1(f"{r['hash']}|{n + 1}".encode()).hexdigest()[:20]


def row_hash(d: date, desc: str, amount: float, balance: str) -> str:
    return hashlib.sha1(f"{d}|{normalize(desc)}|{amount:.2f}|{balance}".encode()).hexdigest()[:20]


def parse(rows: list[list[str]], mapping: dict, header_row: int, owners: list[str] | None = None) -> list[dict]:
    out = []
    for row in rows[header_row + 1:]:
        get = lambda k: row[mapping[k]] if mapping.get(k) is not None and mapping[k] < len(row) else ""  # noqa: E731
        d = parse_date(get("date"))
        desc = get("description")
        if mapping.get("amount") is not None:
            amt = parse_amount(get("amount"))
        else:
            debit, credit = parse_amount(get("debit")) or 0, parse_amount(get("credit")) or 0
            amt = (abs(credit) - abs(debit)) if (debit or credit) else None
        if not d or amt is None or amt == 0 or not desc:
            continue
        out.append({"date": d.isoformat(), "description": desc[:500], "amount": round(amt, 2),
                    "type": classify(desc, amt, owners), "hash": row_hash(d, desc, amt, get("balance")),
                    "balance": parse_amount(get("balance")) if get("balance") else None})
    pair_cancellations(out)
    return out


# ---------------------------------------------------------------- karşı taraf (kişi) çıkarma

_BANK_WORDS = re.compile(r"\b(TURKIYE|T\.?|IS|HALK|ZIRAAT|VAKIFLAR|VAKIFBANK|VAKIF|KUVEYT|TURK|KATILIM|GARANTI|BBVA|YAPI|VE|KREDI|"
                         r"AKBANK|QNB|FINANSBANK|DENIZBANK|ING|TEB|ALBARAKA|ENPARA|PAPARA|EMLAK|CUMHURIYETI|BANKASI|BANKA|"
                         r"BANK|A\.?S\.?|T\.?A\.?S\.?|T\.?A\.?O\.?)\b")


def _clean_name(raw: str) -> str:
    n = normalize(raw)
    n = re.sub(r"TR\d{2}[\d\s*]+(NOLU|NO'LU|NOLU)?", " ", n)
    n = _BANK_WORDS.sub(" ", n)
    n = re.sub(r"\b(ALINAN|GELEN|GIDEN|GONDERILEN|HAVALE|FAST|EFT|ANLIK|ODEME|NOLU|NO LU)\b", " ", n)
    n = re.sub(r"[^A-Z ]", " ", n)
    words = [w for w in n.split() if len(w) > 1]
    cut = next((i for i, w in enumerate(words) if w in _COMPANY_CUT), None)
    if cut:  # şirket: 'ELIF ORNEK EGITIM DANISMANLIK LIMITED SIRKETI' → unvanın başı
        return " ".join(words[max(0, cut - 4):cut][:3])
    return " ".join(words[-3:])


_COMPANY_CUT = {"LIMITED", "LTD", "SIRKETI", "STI", "ANONIM", "TIC", "TICARET", "SAN", "SANAYI", "CORP", "INC", "LLC", "GMBH"}


def _tr_title(word: str, dotless_i: bool = True) -> str:
    low = word.replace("I", "ı" if dotless_i else "i").replace("İ", "i").lower()
    return (low[0].replace("i", "İ").replace("ı", "I").upper() + low[1:]) if low else low


def _pretty(clean: str, original: str) -> str:
    """'ZEYNEP OZDEMIR' → asıl metindeki yazılışla 'Zeynep Özdemir'."""
    tokens = re.findall(r"[^\W\d_]+", original)
    out = []
    for w in clean.split():
        found = [t for t in tokens if normalize(t) == w]
        # büyük harfli yazımda İ/I ayrımı kaybolabiliyor ('DEMIR'): küçük harf içeren yazılışı tercih et
        orig = next((t for t in found if not t.isupper()), found[0] if found else w)
        # metinde hiç 'İ' yoksa banka büyük harfte İ yerine I basıyordur ('AYŞE DEMIR')
        out.append(_tr_title(orig, dotless_i="İ" in original))
    return " ".join(out)


def counterparty(desc: str, amount: float, owners: list[str] | None = None) -> str | None:
    """Havale/FAST açıklamasından karşı tarafın adı. Kendisi ise None."""
    n = normalize(desc)
    own = {_squash(o) for o in owners or [] if o}
    name = None
    sr = _sender_receiver(desc)
    if sr and len(sr[0]) >= 6 and sr[0] == sr[1]:
        return None  # gönderen = alıcı: kendi hesapları arası
    parties = _parties(desc)
    m = re.search(r"GONDEREN\s*:\s*([A-Z ]+?)\s*,\s*ALICI\s*:\s*([A-Z ]+?)\s*(,|$)", n)
    if m:
        name = m.group(2) if amount < 0 else m.group(1)
    elif parties:  # Halkbank/VakıfBank kalıpları: giden → alıcı, gelen → gönderen
        name = parties[1] if amount < 0 else parties[0]
    elif amount < 0 and (m := re.search(r"HESABINDAN\s+(.+?)\s+HESABINA", n)):
        name = m.group(1)
    elif amount > 0 and (m := re.search(r"(?:FAST\s*-\s*)?([A-Z ]{5,}?)\s+TARAFINDAN GONDERILEN", n)):
        name = m.group(1)
    elif amount > 0 and (m := re.search(r"([A-Z ]{5,}?)\s+HESABINDAN", re.sub(r"TR\d{2}[\d\s*]+NOLU|NO'?\s?LU\b", " ", n))):
        name = m.group(1)
    if not name and amount < 0:
        # Halkbank: "GIDEN FAST - X TARAFINDAN GONDERILEN,Alıcı,Banka A.S.,TR.." ya da "not,Alıcı,Banka,TR.."
        parts = [p.strip() for p in re.sub(r"^.*?TARAFINDAN GONDERILEN\s*,", "", n).split(",")]
        bank_i = next((i for i, p in enumerate(parts) if re.search(r"BANK|T\.A\.[OS]|KATILIM", p)), None)
        # yalnızca "…,Alıcı,Banka,TRxx" kalıbı (fatura açıklamalarındaki virgüller karışmasın)
        if bank_i and bank_i + 1 < len(parts) and re.match(r"TR\d{2}", parts[bank_i + 1]):
            name = parts[bank_i - 1]
    if not name:
        return None
    clean = _clean_name(name)
    if not clean or _squash(clean) in own or any(_squash(clean) in o or o in _squash(clean) for o in own):
        return None
    return _pretty(clean, desc)


def transfer_note(desc: str, owners: list[str] | None = None, who: str | None = None) -> str | None:
    """'FAST Anlık Ödeme ev kirası (… sorgu no'lu …)' → 'ev kirası'; Halkbank 'Altin alindi,Alıcı,Banka,TR..' → 'Altin
    alindi'. Not kullanıcının kendi adı ya da karşı tarafın adıysa (banka otomatik doldurur) yok sayılır."""
    m = re.match(r"\s*(?:FAST ANLIK ODEME|GONDERILEN HAVALE|FAST ANLIK ÖDEME|GÖNDERİLEN HAVALE)\s+([^(/\s][^(/]*?)\s*(\(|/|$)",
                 desc, re.I)
    note = m.group(1).strip() if m else None
    if not m and not re.search(r"TARAFINDAN|GONDEREN|HESABINDAN", normalize(desc)):
        # Halkbank: "not,Alıcı,Banka A.S.,TRxx" ya da "not,Gönderen,TRxx,Alıcı,TRxx" — ilk alan kullanıcının notu
        parts = [p.strip() for p in desc.split(",")]
        if len(parts) >= 3 and any(re.match(r"TR\d{2}", p) for p in parts[2:]):
            note = parts[0]
    if not note or re.match(r"TR\d", note) or not re.search(r"[^\W\d_]{2}", note):
        return None
    sq = _squash(note)
    if any(o and len(_squash(o)) >= 6 and _squash(o) == sq for o in [*(owners or []), who or ""]):
        return None
    return note


_SHORT_BANKS = (("HALK", "Halkbank"), ("VAKIF", "VakıfBank"), ("KUVEYT", "Kuveyt Türk"), ("ZIRAAT", "Ziraat"),
                ("GARANTI", "Garanti BBVA"), ("AKBANK", "Akbank"), ("YAPI VE KREDI", "Yapı Kredi"), ("YAPI KREDI", "Yapı Kredi"),
                ("IS BANK", "İş Bankası"), ("ISBANK", "İş Bankası"), ("QNB", "QNB"), ("FINANSBANK", "QNB"),
                ("DENIZ", "DenizBank"), ("ENPARA", "Enpara"), ("ING", "ING"), ("TEB", "TEB"), ("ALBARAKA", "Albaraka"),
                ("EMLAK", "Emlak Katılım"), ("ZIRAAT KATILIM", "Ziraat Katılım"), ("VAKIF KATILIM", "Vakıf Katılım"),
                ("PAPARA", "Papara"), ("ODEABANK", "Odeabank"), ("FIBABANKA", "Fibabanka"), ("SEKERBANK", "Şekerbank"))


def short_bank(name: str | None) -> str:
    """'Türkiye Halk Bankası A.Ş' → 'Halkbank', 'KUVEYT TÜRK KATILIM BANKASI A.Ş.' → 'Kuveyt Türk'."""
    n = normalize(name or "")
    for key, short in sorted(_SHORT_BANKS, key=lambda kv: -len(kv[0])):  # en uzun eşleşme (Ziraat Katılım > Ziraat)
        if re.search(rf"\b{key}", n):
            return short
    return (name or "").strip() or "Banka hesabı"


# ---------------------------------------------------------------- fatura tahsilatı (dökümdeki "X Tahsilatı" satırı)

# (açıklamanın başı, kurum, tür) — sıra önemli: özel olan önce
_BILLS = (
    (r"VODAFONE", "Vodafone", "Telefon faturası"), (r"TURKCELL", "Turkcell", "Telefon faturası"),
    (r"TURK TELEKOM MOBIL|AVEA", "Türk Telekom", "Telefon faturası"),
    (r"TTNET|TURK TELEKOM", "Türk Telekom", "İnternet faturası"), (r"SUPERONLINE", "Superonline", "İnternet faturası"),
    (r"TURKNET", "TurkNet", "İnternet faturası"), (r"DIGITURK", "Digiturk", "TV aboneliği"),
    (r"ENERJISA", "Enerjisa", "Elektrik faturası"), (r"CK BOGAZICI|BOGAZICI ELEKTRIK", "CK Boğaziçi", "Elektrik faturası"),
    (r"AYEDAS", "AYEDAŞ", "Elektrik faturası"), (r"BEDAS", "BEDAŞ", "Elektrik faturası"),
    (r"GEDIZ", "Gediz Elektrik", "Elektrik faturası"), (r"[A-Z]+ ELEKTRIK", None, "Elektrik faturası"),
    (r"IGDAS", "İGDAŞ", "Doğalgaz faturası"), (r"BASKENTGAZ|BASKENT DOGALGAZ", "Başkentgaz", "Doğalgaz faturası"),
    (r"IZMIRGAZ", "İzmirgaz", "Doğalgaz faturası"), (r"[A-Z]+ ?GAZ\b|DOGALGAZ", None, "Doğalgaz faturası"),
    (r"ISKI", "İSKİ", "Su faturası"), (r"ASKI", "ASKİ", "Su faturası"), (r"IZSU", "İZSU", "Su faturası"),
    (r"BUSKI", "BUSKİ", "Su faturası"),
)
_BILL_REFS = (
    ("Telefon", r"TELEFON NUMARASI\s*:\s*(\d{7,})"), ("Abone no", r"(\d{6,})\s+ABONE NOLU"),
    ("Abone no", r"ABONE (?:NO|NUMARASI)\s*:\s*(\d{5,})"), ("Sözleşme no", r"SOZLESME (?:HESAP )?(?:NO|NUMARASI)\s*:\s*(\d{5,})"),
    ("Tesisat no", r"TESISAT (?:NO|NUMARASI)\s*:\s*(\d{5,})"), ("Mukavele no", r"MUKAVELE NO\s*:\s*(\d{5,})"),
)


def bill_info(desc: str) -> dict | None:
    """'Vodafone Tahsilatı Telefon Numarası:54…, Fatura No: FD…, Son Ödeme Tarihi:30.06.2026' →
    {provider: 'Vodafone', kind: 'Telefon faturası', ref: 'Telefon 54…', invoice: 'FD…', due: '2026-06-30'}.
    Yalnız fatura tahsilatı satırları (açıklama 'X Tahsilatı' ile başlar ya da 'Fatura No' içerir); masraf/vergi değil."""
    n = normalize(desc)
    head = re.match(r"\s*(.{2,60}?)\s+(?:FATURA\s+)?TAHSILATI?\b", n)
    if not head and "FATURA NO" not in n:
        return None
    name = head.group(1) if head else n[:60]
    if re.search(r"\b(MASRAF|VERGI|SADAKA|BAGIS|KOMISYON|HARC|SURUCU|TRAFIK|CEZA|SGK|KREDI|AIDAT)\b", name):
        return None
    provider = kind = None
    for pat, prov, k in _BILLS:
        if re.search(pat, name):
            provider, kind = prov, k
            break
    if not kind:
        return None
    if not provider:  # 'Bursa Elektrik Tahsilatı' → asıl yazılışla kurum adı
        words = re.findall(r"[^\W\d_]+", desc)[: len(name.split())]
        provider = " ".join(w if w.isupper() and len(w) <= 5 else w.capitalize() for w in words) or name.title()
    ref = next((f"{label} {m.group(1)}" for label, pat in _BILL_REFS if (m := re.search(pat, n))), None)
    inv = re.search(r"FATURA (?:NO|NUMARASI)\s*:\s*([A-Z0-9]{6,})", n)
    due = re.search(r"SON ODEME TARIHI\s*:\s*(\d{2})[./](\d{2})[./](\d{4})|SON ODEME TARIHI\s*:\s*(\d{4})/(\d{2})/(\d{2})", n)
    due_iso = None
    if due:
        g = due.groups()
        due_iso = f"{g[2]}-{g[1]}-{g[0]}" if g[0] else f"{g[3]}-{g[4]}-{g[5]}"
    return {"provider": provider, "kind": kind, "ref": ref, "invoice": inv.group(1) if inv else None, "due": due_iso}


def describe(raw: str, amount: float, owners: list[str] | None = None) -> dict:
    """Dökümdeki ham açıklamadan kayıt alanları: {description, counterparty, payment, note, bill_provider}.
    İçe aktarma ve onarım (app/bank_repair.py) aynı sonucu üretir. Kişi bulunamazsa ve fatura değilse açıklama ham kalır
    (uzunsa 120 karakter; tamamı notta)."""
    who = counterparty(raw, amount, owners)
    norm = normalize(raw)
    payment = "havale" if (who or re.search(r"EFT|HAVALE|FAST|TRANSFER", norm)) else "nakit"
    desc, extra, provider = raw, [], None
    if who:
        note = transfer_note(raw, owners, who)
        desc = who + (f" · {note}" if note else "")
    elif amount < 0 and (b := bill_info(raw)):
        provider = b["provider"]
        desc = f"{b['provider']} · {b['kind']}"
        extra = [x for x in (b["ref"], b["invoice"] and f"Fatura no {b['invoice']}",
                             b["due"] and f"Son ödeme {b['due'][8:10]}.{b['due'][5:7]}.{b['due'][:4]}") if x]
    elif amount < 0 and re.match(r"\s*MASRAF TAHSILAT", norm):
        payment = "nakit"
        desc = "Banka masrafı" + (" · ATM komisyonu" if "ATM" in norm else "")
    elif amount < 0 and re.search(r"PARA CEK|\bPR CKM\b", norm):  # ATM'den nakit
        payment = "nakit"
        branch = re.search(r"NO LU (.+?) ATM", norm)
        where = (_pretty_words(branch.group(1), raw) if branch and re.search(r"[A-Z]{3}", branch.group(1))
                 else "diğer banka" if "DIGER BANKA" in norm
                 else "ortak nokta" if "ORTAK NOKTA" in norm else None)  # şube yerine terminal no yazıyorsa atla
        desc = ("Para çekme · şube" if "MEVDUAT PARA CEK" in norm else "ATM para çekme" + (f" · {where}" if where else ""))
    if desc == raw and len(desc) > 120:
        desc = desc[:120]
    note = "Hesap hareketi" + "".join(f" · {x}" for x in extra) + ("" if desc == raw else f" · {raw}")
    return {"description": desc, "counterparty": who, "payment": payment, "note": note, "bill_provider": provider}


def _pretty_words(clean: str, original: str) -> str:
    """'MASLAK TICARI SUBESI' → asıl metindeki yazılışla 'Maslak Ticari Şubesi'."""
    tokens = re.findall(r"[^\W\d_]+", original)
    out = []
    for w in clean.split():
        orig = next((t for t in tokens if normalize(t) == w), w)
        out.append(_tr_title(orig, dotless_i="İ" in original) if orig.isupper() else orig)
    return " ".join(out)


def is_cancellation(desc: str) -> bool:
    """'FAST Anlık Ödeme İptali … hesaba geri ödemesi' — giden ödemenin bankaca iadesi."""
    return bool(re.search(r"\bIPTAL(I|EN)?\b|GERI ODEMES", normalize(desc)))


def pair_cancellations(rows: list[dict]) -> None:
    """Aynı dökümde iptal edilen ödeme ile iadesi: ikisi de 'iptal' (gelir/gider sayılmaz, kayıt açılmaz). Eşleşme:
    aynı tutar, ±2 gün, iade satırı iptal/geri ödeme diyor; sorgu/referans numarası varsa aynı olmalı."""
    used: set[int] = set()
    for i, r in enumerate(rows):
        if r.get("amount", 0) <= 0 or not is_cancellation(r["description"]):
            continue
        refs = set(re.findall(r"\b\d{9,12}\b", r["description"]))
        d = parse_date(r["date"])
        # en yakın satırdan başla (önce öncekiler): aynı gün aynı tutarda başka çekim de olabilir
        for j in sorted((k for k in range(len(rows)) if k != i), key=lambda k: (abs(k - i), k > i)):
            o = rows[j]
            if j in used or o.get("type") not in ("gider", "virman") or abs(o["amount"] + r["amount"]) > 0.005:
                continue
            od = parse_date(o["date"])
            if not d or not od or abs((d - od).days) > 2:
                continue
            orefs = set(re.findall(r"\b\d{9,12}\b", o["description"]))
            if refs and orefs and not refs & orefs:
                continue
            r["type"] = o["type"] = "iptal"
            used |= {i, j}
            break


def _bank_in(text: str) -> str | None:
    """Metindeki banka adı ya da IBAN'ın banka kodundan kısa banka adı."""
    from app.templates import IBAN_BANKS

    m = re.search(r"TR\s?\d{2}\s?(\d{4})\s?(\d)", text.upper())
    if m and (code := m.group(1) + m.group(2)) in IBAN_BANKS:
        return IBAN_BANKS[code]
    b = short_bank(text)
    return None if b in ("Banka hesabı", (text or "").strip()) else b


def own_label(raw: str, amount: float, account: str | None, owners: list[str] | None = None) -> str:
    """Kendi hesapları arası aktarımın temiz açıklaması: 'Kendi hesabına · Halkbank → VakıfBank · not'. Gönderenin
    yazdığı mesaj (not) kaybolmaz; ham banka metni own_transfers.raw'da kalır."""
    acc = short_bank(account) if account else None
    other = None
    p = _parties(raw)
    if p:
        other = _bank_in(p[1] if amount < 0 else p[0])
    if not other or other == acc:
        from app.templates import IBAN_BANKS

        n = normalize(raw)
        tail = re.split(r"TARAFINDAN GONDERILEN", n, maxsplit=1)[-1]
        found = [IBAN_BANKS.get(a + b) for a, b in re.findall(r"TR\s?\d{2}\s?(\d{4})\s?(\d)", tail)] + [_bank_in(tail)]
        other = next((b for b in found if b and b != acc), None)
    if amount < 0:
        label = "Kendi hesabına" + (f" · {acc or '?'} → {other}" if other else (f" · {acc}" if acc else ""))
    else:
        label = "Kendi hesabından" + (f" · {other} → {acc or '?'}" if other else (f" · {acc}" if acc else ""))
    note = transfer_note(raw, owners)
    return label + (f" · {note}" if note else "")
